import io

from openpyxl import load_workbook
from PyPDF2 import PdfReader

from odoo.exceptions import UserError
from odoo.tests import tagged

from .common import JANUARY, S03aCommon


@tagged('post_install', '-at_install')
class TestS03aExports(S03aCommon):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.sale = cls._post('2026-01-05', [
            ('131', 110, 0, {'name': "Sale"}),
            ('511', 0, 100, {'name': "Sale"}),
            ('33311', 0, 10, {'name': "VAT"}),
        ])

    def _many_entries(self, count=120):
        for index in range(count):
            self._post('2026-01-15', [
                ('642', 10 + index, 0, {'name': f"Expense {index}"}),
                (self.cash, 0, 10 + index, {'name': f"Expense {index}"}),
            ])

    def _pdf_pages(self, **options):
        content = self.Report.get_pdf_rl({**JANUARY, **options})
        self.assertTrue(content.startswith(b'%PDF'))
        return [' '.join(page.extract_text().split()) for page in PdfReader(io.BytesIO(content)).pages]

    # ------------------------------------------------------------------ Excel

    def test_xlsx(self):
        content = self.Report.get_xlsx(dict(JANUARY))
        rows = list(load_workbook(io.BytesIO(content))['Sheet1'].iter_rows(values_only=True))
        self.assertEqual(rows[0][0], "GENERAL JOURNAL")
        self.assertEqual(rows[1][0], "From 2026-01-01 to 2026-01-31")
        self.assertEqual(rows[3], (
            "Date", "Voucher No.", "Description", "Debit Account", "Credit Account", "Debit", "Credit",
        ))
        self.assertEqual(rows[4], ('2026/01/05', self.sale.name, "Sale", '131', None, 110, None))
        self.assertEqual(rows[5], ('2026/01/05', self.sale.name, "Sale", None, '511', None, 100))
        self.assertEqual(rows[-1], (None, None, "Total", None, None, 110, 110))
        self.assertEqual(len(rows), 4 + 3 + 1)

    def test_xlsx_exports_the_whole_journal_when_the_screen_is_lazy(self):
        self._many_entries(20)
        self.env['ir.config_parameter'].sudo().set_int('fg_vn_accounting_reports.lazy_line_threshold', 5)
        self.env['ir.config_parameter'].sudo().set_int('fg_vn_accounting_reports.lazy_chunk_size', 4)
        self.assertTrue(self.Report.get_report_data(dict(JANUARY))['lazy'])
        content = self.Report.get_xlsx(dict(JANUARY))
        rows = list(load_workbook(io.BytesIO(content))['Sheet1'].iter_rows(values_only=True))
        self.assertEqual(len(rows), 4 + 3 + 40 + 1)

    # -------------------------------------------------------------------- PDF

    def test_pdf_single_page(self):
        pages = self._pdf_pages()
        self.assertEqual(len(pages), 1)
        text = pages[0]
        for label in (
            "GENERAL JOURNAL", "Voucher No.", self.sale.name, "33311", "Total",
            "This book has 1 pages, numbered from page 1 to page 1",
            "Book opening date: 2026-01-01", "Prepared by", "Page 1/1",
        ):
            self.assertIn(label, text)
        self.assertNotIn("Carried forward", text)

    def test_pdf_carries_the_totals_from_page_to_page(self):
        self._many_entries()
        pages = self._pdf_pages()
        self.assertGreater(len(pages), 2)
        self.assertIn("Carried forward", pages[0])
        self.assertIn("Brought forward", pages[1])
        self.assertNotIn("Carried forward", pages[-1])
        for text in pages:
            self.assertIn("GENERAL JOURNAL", text)
            self.assertIn("Voucher No.", text)
        self.assertIn(f"This book has {len(pages)} pages", pages[-1])

    def test_pdf_orientation(self):
        for orientation, portrait in (('portrait', True), ('landscape', False)):
            content = self.Report.get_pdf_rl({**JANUARY, 'orientation': orientation})
            box = PdfReader(io.BytesIO(content)).pages[0].mediabox
            self.assertEqual(float(box[2]) < float(box[3]), portrait)

    def test_pdf_in_vietnamese(self):
        self.env['res.lang']._activate_lang('vi_VN')
        self.env['ir.module.module']._load_module_terms(
            ['fg_vn_accounting_reports', 'fg_vn_report_s03a'], ['vi_VN'],
        )
        self._many_entries()
        pages = self._pdf_pages(lang='vi_VN')
        for label in ("SỔ NHẬT KÝ CHUNG", "Số chứng từ", "Diễn giải", "TK Nợ", "TK Có", "Cộng chuyển trang"):
            self.assertIn(label, pages[0])
        self.assertIn("Số dư đầu trang", pages[1])
        self.assertIn("Người lập biểu", pages[-1])
        data = self.Report.get_report_data({'date': JANUARY['date']})
        self.assertEqual(data['header']['template_code'], "Mẫu số S03a-DNN")
        self.assertEqual(data['lines'][0]['total_label'], "Cộng")

    def test_pdf_line_limit(self):
        self._many_entries(10)
        self.env['ir.config_parameter'].sudo().set_int('fg_vn_accounting_reports.pdf_line_limit', 20)
        with self.assertRaises(UserError):
            self.Report.get_pdf_rl(dict(JANUARY))
        # narrowing the journal makes it printable again
        self.assertTrue(self.Report.get_pdf_rl({**JANUARY, 'account_ids': self.accounts['131'].ids}))
