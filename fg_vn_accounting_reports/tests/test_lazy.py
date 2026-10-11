import io

from openpyxl import load_workbook

from odoo.exceptions import AccessError
from odoo.tests import TransactionCase, tagged

from .common import REPORT_MODEL, FakeReportMixin, make_groups

OPTIONS = {'date': {'date_from': '2026-01-01', 'date_to': '2026-01-31'}, 'lang': 'en_US'}


@tagged('post_install', '-at_install')
class TestLazyReport(FakeReportMixin, TransactionCase):
    """ Long reports: over the threshold, the screen gets the groups without
    their lines and loads the lines by chunks. """

    def setUp(self):
        super().setUp()
        self.groups = make_groups(group_count=2, line_count=5)  # 10 lines
        self.countable = True
        self._set_param('lazy_line_threshold', 8)
        self._set_param('lazy_chunk_size', 2)

    def _set_param(self, name, value):
        self.env['ir.config_parameter'].sudo().set_int(f'fg_vn_accounting_reports.{name}', value)

    def test_report_that_cannot_count_is_never_lazy(self):
        self.countable = False
        data = self.Report.get_report_data(dict(OPTIONS))
        self.assertFalse(data['lazy'])
        self.assertEqual(len(data['lines'][0]['lines']), 5)

    def test_under_the_threshold_everything_is_sent(self):
        self._set_param('lazy_line_threshold', 10)
        data = self.Report.get_report_data(dict(OPTIONS))
        self.assertFalse(data['lazy'])
        self.assertEqual([len(group['lines']) for group in data['lines']], [5, 5])

    def test_over_the_threshold_groups_come_without_lines(self):
        data = self.Report.get_report_data(dict(OPTIONS))
        self.assertTrue(data['lazy'])
        self.assertEqual(data['chunk_size'], 2)
        self.assertEqual([group['lines'] for group in data['lines']], [[], []])
        self.assertEqual([group['line_count'] for group in data['lines']], [5, 5])
        self.assertEqual(data['lines'][0]['name'], "Account 110")
        self.assertNotIn('preload', data['lines'][0])

    def test_preloaded_group_comes_with_its_first_chunk(self):
        self.preload = True
        data = self.Report.get_report_data(dict(OPTIONS))
        self.assertEqual(
            [line['id'] for line in data['lines'][0]['lines']], ['line_0_0', 'line_0_1'],
        )

    def test_group_lines_are_loaded_by_chunks(self):
        loaded = []
        offset = 0
        while True:
            result = self.Report.get_group_lines(dict(OPTIONS), 'group_1', offset)
            if not result['lines']:
                break
            loaded += result['lines']
            offset = result['offset']
        self.assertEqual([line['id'] for line in loaded], [f'line_1_{i}' for i in range(5)])
        self.assertEqual(offset, 5)

    def test_group_lines_check_the_access(self):
        Report = self.env[REPORT_MODEL].with_user(self.other_user)
        with self.assertRaises(AccessError):
            Report.get_group_lines(dict(OPTIONS), 'group_0', 0)
        other_company = self.env['res.company'].create({'name': "FG other company"})
        with self.assertRaises(AccessError):
            self.Report.get_group_lines({**OPTIONS, 'company_id': other_company.id}, 'group_0', 0)

    def test_group_lines_are_in_the_report_language(self):
        seen = []
        self.patch(
            self.registry[REPORT_MODEL], '_get_group_lines',
            lambda report, options, group_id, offset, limit: seen.append(report.env.lang) or [],
        )
        self.Report.get_group_lines({'date': OPTIONS['date']}, 'group_0', 0)
        self.Report.get_group_lines(dict(OPTIONS), 'group_0', 0)
        self.assertEqual(seen, ['vi_VN', 'en_US'])

    def test_exports_are_never_lazy(self):
        content = self.Report.get_xlsx(dict(OPTIONS))
        rows = list(load_workbook(io.BytesIO(content))['Sheet1'].iter_rows(values_only=True))
        # 4 heading rows; a header, 5 lines and a total for each of the 2 groups
        self.assertEqual(len(rows), 4 + 2 * 7)
        self.assertTrue(self.Report.get_pdf_rl(dict(OPTIONS)).startswith(b'%PDF'))

    def test_invalid_parameters_fall_back_on_the_defaults(self):
        self.env['ir.config_parameter'].sudo().set_str('fg_vn_accounting_reports.lazy_line_threshold', 'abc')
        self._set_param('lazy_chunk_size', 0)
        self.assertEqual(self.Report._get_lazy_threshold(), 2000)
        self.assertEqual(self.Report._get_lazy_chunk_size(), 1)
