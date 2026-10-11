import io

from openpyxl import load_workbook

from odoo.tests import TransactionCase, tagged

from .common import FakeReportMixin

OPTIONS = {'date': {'date_from': '2026-01-01', 'date_to': '2026-01-31'}, 'lang': 'en_US'}


@tagged('post_install', '-at_install')
class TestReportXlsx(FakeReportMixin, TransactionCase):

    def _load(self, options=None):
        content = self.Report.get_xlsx(dict(options or OPTIONS))
        return load_workbook(io.BytesIO(content))

    def test_single_sheet(self):
        workbook = self._load()
        self.assertEqual(workbook.sheetnames, ['Sheet1'])
        rows = list(workbook['Sheet1'].iter_rows(values_only=True))

        self.assertEqual(rows[0][0], "Report")
        self.assertEqual(rows[1][0], "From 2026-01-01 to 2026-01-31")
        self.assertEqual(rows[3], ("Date", "Label", "Debit", "Credit"))
        # group header, 3 lines and a total for each of the 2 groups
        self.assertEqual(len(rows), 4 + 2 * 5)
        self.assertEqual(rows[4][0], "Account 110")
        self.assertEqual(rows[5], ('2026-01-01', "Entry 0/0", 1000, None))
        self.assertEqual(rows[8][1:3], ("Total", 6000))

    def test_amounts_are_numbers_and_zero_is_blank(self):
        sheet = self._load()['Sheet1']
        debit, credit = sheet['C6'], sheet['D6']
        self.assertEqual(debit.data_type, 'n')
        self.assertEqual(debit.number_format, '#,##0')
        self.assertIsNone(credit.value)

    def test_group_header_spans_all_columns(self):
        sheet = self._load()['Sheet1']
        self.assertIn('A5:D5', [str(merged) for merged in sheet.merged_cells.ranges])

    def test_one_sheet_per_group(self):
        self.split_by_group = True
        workbook = self._load()
        self.assertEqual(workbook.sheetnames, ["Account 110", "Account 111"])
        rows = list(workbook["Account 111"].iter_rows(values_only=True))
        self.assertEqual(rows[0][0], "Report")
        self.assertEqual(rows[3], ("Date", "Label", "Debit", "Credit"))
        self.assertEqual(len(rows), 4 + 5)
        self.assertEqual(rows[5][1], "Entry 1/0")

    def test_sheet_names_are_valid_and_unique(self):
        self.split_by_group = True
        self.groups[0]['name'] = "131/Customers: [VND] with a very long account name"
        self.groups[1]['name'] = "131/Customers: [VND] with a very long account name"
        names = self._load().sheetnames
        self.assertEqual(len(set(names)), 2)
        for name in names:
            self.assertLessEqual(len(name), 31)
            self.assertFalse(set(name) & set('/\\*?[]:'))

    def test_without_figures(self):
        self.groups = []
        rows = list(self._load()['Sheet1'].iter_rows(values_only=True))
        self.assertEqual(len(rows), 4)

    def test_period_line_follows_the_report_language(self):
        self.env['ir.module.module']._load_module_terms(['fg_vn_accounting_reports'], ['vi_VN'])
        rows = list(self._load({**OPTIONS, 'lang': 'vi_VN'})['Sheet1'].iter_rows(values_only=True))
        self.assertEqual(rows[1][0], "Từ ngày 2026-01-01 đến ngày 2026-01-31")
