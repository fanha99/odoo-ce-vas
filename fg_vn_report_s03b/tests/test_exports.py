import io

from openpyxl import load_workbook
from PyPDF2 import PdfReader

from odoo.tests import tagged

from .common import JANUARY, S03bCommon


@tagged('post_install', '-at_install')
class TestS03bExports(S03bCommon):

    def _many_items(self, count=110):
        for index in range(count):
            self._post('2026-01-22', [
                ('131', 10 + index, 0, {'name': f"Sale {index}", 'partner_id': self.partner_a.id}),
                ('511', 0, 10 + index, {'name': f"Sale {index}"}),
            ])

    def _pdf_pages(self, **options):
        content = self.Report.get_pdf_rl({**JANUARY, **options})
        self.assertTrue(content.startswith(b'%PDF'))
        return [' '.join(page.extract_text().split()) for page in PdfReader(io.BytesIO(content)).pages]

    # ------------------------------------------------------------------ Excel

    def _workbook(self, **options):
        return load_workbook(io.BytesIO(self.Report.get_xlsx({**JANUARY, **options})))

    def test_xlsx_has_one_sheet_per_account(self):
        workbook = self._workbook()
        groups = self._groups()
        self.assertEqual(len(workbook.sheetnames), len(groups))
        self.assertIn("131 - Test 131", workbook.sheetnames)
        for name in workbook.sheetnames:
            self.assertLessEqual(len(name), 31)

    def test_xlsx_sheet_of_an_account(self):
        rows = list(self._workbook()["131 - Test 131"].iter_rows(values_only=True))
        self.assertEqual(rows[0][0], "GENERAL LEDGER")
        self.assertEqual(rows[1][0], "From 2026-01-01 to 2026-01-31")
        self.assertEqual(rows[2][0], "Account: 131 - Test 131")
        self.assertEqual(rows[4], (
            "Date", "Voucher No.", "Description", "Counterpart Account", "Debit", "Credit",
        ))
        self.assertEqual(rows[5], (None, None, "Opening balance", None, 200, None))
        self.assertEqual(rows[6], ('2026/01/05', self.sale.name, "Sale", '33311, 511', 110, None))
        self.assertEqual(rows[7][3:], (self.bank.code, None, 110))
        self.assertEqual(rows[8], (None, None, "Total of the period", None, 110, 110))
        self.assertEqual(rows[9], (None, None, "Closing balance", None, 200, None))
        self.assertEqual(len(rows), 10)

    def test_xlsx_split_counterparts(self):
        rows = list(self._workbook(split_counterparts=True)["131 - Test 131"].iter_rows(values_only=True))
        self.assertEqual([row[3:5] for row in rows[6:8]], [('33311', 10), ('511', 100)])
        self.assertEqual(rows[9], (None, None, "Total of the period", None, 110, 110))

    # -------------------------------------------------------------------- PDF

    def test_pdf_starts_each_account_on_a_new_page(self):
        pages = self._pdf_pages()
        groups = self._groups()
        self.assertEqual(len(pages), len(groups))
        for text, group in zip(pages, groups.values()):
            self.assertIn("GENERAL LEDGER", text)
            self.assertIn(group['label'], text, "the account is named in the page header")
            for label in ("Opening balance", "Total of the period", "Closing balance", "Prepared by"):
                self.assertIn(label, text)
            self.assertIn("This book has 1 pages", text)

    def test_pdf_carries_the_totals_within_an_account(self):
        self._many_items()
        pages = self._pdf_pages(account_ids=(self.accounts['131'] + self.accounts['642']).ids)
        receivable = [text for text in pages if "Account: 131 - Test 131" in text]
        expense = [text for text in pages if "Account: 642 - Test 642" in text]
        self.assertGreater(len(receivable), 1)
        self.assertEqual(len(expense), 1)
        self.assertEqual(len(receivable) + len(expense), len(pages))
        self.assertIn("Carried forward", receivable[0])
        self.assertIn("Brought forward", receivable[1])
        self.assertNotIn("Brought forward", expense[0], "the totals restart with each account")
        # the total of the period and the closing balance stay together
        last = receivable[-1]
        self.assertIn("Total of the period", last)
        self.assertIn("Closing balance", last)
        self.assertIn(f"This book has {len(receivable)} pages", last)

    def test_pdf_mirror_margins_start_accounts_on_odd_pages(self):
        pages = self._pdf_pages(
            account_ids=(self.accounts['131'] + self.accounts['642']).ids, mirror_margin=True,
        )
        self.assertEqual(len(pages), 3)
        self.assertIn("Account: 131", pages[0])
        self.assertEqual(pages[1], '')
        self.assertIn("Account: 642", pages[2])

    def test_pdf_in_vietnamese(self):
        self.env['res.lang']._activate_lang('vi_VN')
        self.env['ir.module.module']._load_module_terms(
            ['fg_vn_accounting_reports', 'fg_vn_report_s03b'], ['vi_VN'],
        )
        pages = self._pdf_pages(account_ids=self.accounts['131'].ids, lang='vi_VN')
        for label in (
            "SỔ CÁI TÀI KHOẢN", "Tài khoản: 131", "TK đối ứng", "Số dư đầu kỳ",
            "Cộng phát sinh", "Số dư cuối kỳ", "Người lập biểu",
        ):
            self.assertIn(label, pages[0])
        data = self.Report.get_report_data({'date': JANUARY['date']})
        self.assertEqual(data['header']['template_code'], "Mẫu số S03b-DNN")
