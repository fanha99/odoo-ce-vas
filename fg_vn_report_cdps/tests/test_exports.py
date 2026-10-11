import io
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from PyPDF2 import PdfReader

from odoo.tests import tagged

from .common import JANUARY, CDPSCommon


@tagged('post_install', '-at_install')
class TestCDPSExports(CDPSCommon):

    # --------------------------------------------------------------- language

    def test_report_in_vietnamese(self):
        """ Without option, the report is in Vietnamese on screen, in Excel
        and in PDF, whatever the language of the user. """
        self.env['res.lang']._activate_lang('vi_VN')
        self.env['ir.module.module']._load_module_terms(
            ['fg_vn_accounting_reports', 'fg_vn_report_cdps'], ['vi_VN'],
        )
        self.assertEqual(self.env.user.lang, 'en_US')
        options = {'date': JANUARY['date'], 'comparison_type': 'previous_period'}

        data = self.Report.get_report_data(dict(options))
        self.assertEqual(data['header']['title'], "BẢNG CÂN ĐỐI SỐ PHÁT SINH")
        self.assertEqual(
            [column['label'] for column in data['columns'][:4]],
            ["Mã TK", "Tên tài khoản", "Dư Nợ đầu kỳ", "Dư Có đầu kỳ"],
        )
        self.assertEqual(data['columns'][10]['label'], "Chênh lệch")
        report = data['lines'][0]
        self.assertEqual(report['total_label'], "TỔNG CỘNG")
        self.assertEqual(report['comparison_label'], "Kỳ trước (01/12/2025 - 31/12/2025)")

        content = self.Report.get_xlsx(dict(options))
        rows = list(load_workbook(io.BytesIO(content))['Sheet1'].iter_rows(values_only=True))
        self.assertEqual(rows[0][0], "BẢNG CÂN ĐỐI SỐ PHÁT SINH")
        self.assertEqual(rows[1][0], "Từ ngày 2026-01-01 đến ngày 2026-01-31")
        self.assertEqual(rows[3][0], "Mã TK")
        self.assertEqual(rows[-1][1], "TỔNG CỘNG")

        content = self.Report.get_pdf_rl(dict(options))
        text = ' '.join(page.extract_text() for page in PdfReader(io.BytesIO(content)).pages)
        text = ' '.join(text.split())
        for label in ("BẢNG CÂN ĐỐI SỐ PHÁT SINH", "Tên tài khoản", "Chênh lệch", "TỔNG CỘNG", "Trang 1/1"):
            self.assertIn(label, text)

    # ------------------------------------------------------------------ Excel

    def _xlsx_rows(self, **options):
        content = self.Report.get_xlsx({**JANUARY, **options})
        return list(load_workbook(io.BytesIO(content))['Sheet1'].iter_rows(values_only=True))

    def test_xlsx(self):
        rows = self._xlsx_rows()
        lines = self._report()['lines']
        self.assertEqual(rows[0][0], "TRIAL BALANCE")
        self.assertEqual(rows[3][:3], ("Account", "Account Name", "Opening Debit"))
        # title, period, blank, column headers, the lines, the total
        self.assertEqual(len(rows), 4 + len(lines) + 1)

        by_code = {row[0]: row for row in rows[4:]}
        self.assertEqual(by_code['1111'][2:], (1000, None, None, 300, 700, None))
        self.assertEqual(rows[-1][1:], ("TOTAL", 1000, 1000, 1120, 1120, 1820, 1820))

    def test_xlsx_hierarchy_indents_the_children(self):
        rows = self._xlsx_rows(hierarchy=True)
        codes = [row[0] for row in rows[4:-1]]
        self.assertIn('111', codes)
        self.assertIn('   1111', codes)

    def test_xlsx_comparison(self):
        rows = self._xlsx_rows(comparison_type='previous_period')
        self.assertEqual(len(rows[3]), 12)
        cash = next(row for row in rows if row[0] == '1111')
        self.assertEqual(cash[8:], (1000, None, -300, '-30.0%'))

    # -------------------------------------------------------------------- PDF

    def _pdf_text(self, **options):
        content = self.Report.get_pdf_rl({**JANUARY, **options})
        self.assertTrue(content.startswith(b'%PDF'))
        return '\n'.join(page.extract_text() for page in PdfReader(io.BytesIO(content)).pages)

    def test_pdf(self):
        text = self._pdf_text()
        self.assertIn("TRIAL BALANCE", text)
        self.assertIn("Test 1111", text)
        self.assertIn("TOTAL", text)
        self.assertIn("1,820", text)
        self.assertNotIn("Prepared by", text, "the trial balance is not signed")
        self.assertNotIn("Carried forward", text)

    def test_pdf_comparison(self):
        text = self._pdf_text(comparison_type='previous_period', hierarchy=True)
        self.assertIn("Variance", text)
        self.assertIn("-30.0%", text)

    def test_pdf_portrait(self):
        content = self.Report.get_pdf_rl({**JANUARY, 'orientation': 'portrait'})
        box = PdfReader(io.BytesIO(content)).pages[0].mediabox
        self.assertLess(float(box[2]), float(box[3]))

    # -------------------------------------------------------------------- XML

    def _xml(self, **options):
        content = self.Report.get_xml({**JANUARY, **options})
        self.assertIsInstance(content, bytes)
        self.assertTrue(content.startswith(b'<?xml version="1.0" encoding="utf-8"?>'))
        return ET.fromstring(content)

    def _xml_value(self, root, section, side, tag):
        return root.find(f'{section}/{side}/{tag}').text

    def test_xml_structure(self):
        root = self._xml()
        self.assertEqual(root.tag, 'PL_CDTK')
        self.assertEqual([node.tag for node in root], ['SoDuDauKy', 'SoPhatSinhTrongKy', 'SoDuCuoiKy'])
        for section in root:
            self.assertEqual([node.tag for node in section], ['No', 'Co'])
            self.assertEqual(section[0][-1].tag, 'tongCong')

    def test_xml_figures(self):
        root = self._xml()
        self.assertEqual(self._xml_value(root, 'SoDuDauKy', 'No', 'ct1111'), '1000')
        self.assertEqual(self._xml_value(root, 'SoPhatSinhTrongKy', 'Co', 'ct1111'), '300')
        self.assertEqual(self._xml_value(root, 'SoDuCuoiKy', 'No', 'ct1111'), '700')
        self.assertEqual(self._xml_value(root, 'SoDuCuoiKy', 'No', 'ct1311'), '500')
        self.assertEqual(self._xml_value(root, 'SoDuCuoiKy', 'Co', 'ct1311'), '200')

    def test_xml_parent_account_sums_its_children(self):
        root = self._xml()
        self.assertEqual(self._xml_value(root, 'SoPhatSinhTrongKy', 'No', 'ct111'), '250')
        self.assertEqual(self._xml_value(root, 'SoDuCuoiKy', 'No', 'ct111'), '950')
        self.assertEqual(self._xml_value(root, 'SoDuCuoiKy', 'Co', 'ct131'), '200')

    def test_xml_lists_the_accounts_without_figures(self):
        root = self._xml()
        self.assertEqual(self._xml_value(root, 'SoDuCuoiKy', 'No', 'ct113'), '0')

    def test_xml_leaves_off_balance_accounts_out(self):
        root = self._xml()
        tags = {node.tag for node in root.find('SoDuCuoiKy/No')}
        self.assertFalse({'ctN01', 'ctN02', 'ct0091', 'ct0092'} & tags)
        # the grand total of the report includes the off-balance accounts (70), the XML does not
        self.assertEqual(self._xml_value(root, 'SoPhatSinhTrongKy', 'No', 'tongCong'), '1050')
        self.assertEqual(self._xml_value(root, 'SoPhatSinhTrongKy', 'Co', 'tongCong'), '1050')
        self.assertEqual(self._xml_value(root, 'SoDuCuoiKy', 'No', 'tongCong'), '1750')
        self.assertEqual(self._xml_value(root, 'SoDuCuoiKy', 'Co', 'tongCong'), '1750')

    def test_xml_ignores_the_filters(self):
        filtered = self.Report.get_xml({
            **JANUARY, 'account_ids': self.accounts['1111'].ids, 'hierarchy': True,
        })
        self.assertEqual(filtered, self.Report.get_xml(dict(JANUARY)))

    def test_xml_rounds_the_amounts(self):
        self._post('2026-01-23', '642', '1111', 0.6)
        root = self._xml()
        self.assertEqual(self._xml_value(root, 'SoPhatSinhTrongKy', 'No', 'ct642'), '301')

    def test_xml_hidden_transfer_account_is_out_of_the_total(self):
        self._post('2026-01-15', '113', '1111', 100)
        self._post('2026-01-16', '1112', '113', 100)
        self.company.write({
            'fg_hide_transfer_account': True,
            'fg_hidden_transfer_account_ids': self.accounts['113'].ids,
        })
        root = self._xml()
        self.assertEqual(self._xml_value(root, 'SoPhatSinhTrongKy', 'No', 'ct113'), '100')
        self.assertEqual(self._xml_value(root, 'SoPhatSinhTrongKy', 'No', 'tongCong'), '1150')
        self.assertEqual(self._xml_value(root, 'SoPhatSinhTrongKy', 'Co', 'tongCong'), '1150')
